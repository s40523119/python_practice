#!/usr/bin/env python
# coding: utf-8

# In[1]:


import os, re
import docx
from docx import Document


# In[2]:


def parse_priority_year(path):
    root = os.getcwd()
    docx_path = root + path
    document = Document(docx_path)
    tables = document.tables
    table = tables[0]
    year = (table.cell(2,4)).text[:4]
    return year


# In[3]:


def parse_all_docx_path():
    files_path = list()
    files_list = list()
    for root, dirs, files in os.walk("."):
        if root[-2:] == '練習' or root[-2:] == '例表':
            continue
        for file in files:
            if file.endswith(".docx"):
                if file[:2] == 'CN' or file[:2] == 'EP' or file[:2] == 'UA':
                    #print(os.path.join(root,file))
                    files_list.append(file)
                    files_path.append((os.path.join(root,file))[1:])
    return files_path,  files_list


# In[4]:


def gPriorityYearList():
    file_path = parse_all_docx_path()[0]
    priority_list = set([(parse_priority_year(path)) for path in file_path])
    arg_filter = re.compile("\d\d\d\d")
    priority_list = list(filter(arg_filter.match, priority_list))
    priority_list.sort()
    return priority_list


# In[5]:


file_path = parse_all_docx_path()[0]
fname = parse_all_docx_path()[1]
result = dict()
for path in file_path:
    patent = fname[file_path.index(path)][:-5]
    year = parse_priority_year(path)
    result.update({patent: year})
    #print(fname[file_path.index(path)][:-5] + ' : ' + parse_priority_year(path))


# In[6]:


result = sorted(result.items(), key=lambda kv: kv[1])
result


# In[7]:





# In[ ]:





# In[ ]:





# In[ ]:




